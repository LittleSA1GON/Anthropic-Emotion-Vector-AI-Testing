"""Create rubric-aligned editable CMPE187 report (DOCX), plus derived PDF.
Never invent results; unexecuted studies receive a prominent Not Yet Run status.
"""
from __future__ import annotations
import argparse
import json
import math
import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
from .course_tests import ABSTRACTIONS, DICTIONS, EMOTIONS, BOUNDS, RAW_RANGES, starter_cases, summary, validate_cases


def render_drawio_png(xml_path, output_dir):
    from PIL import Image, ImageDraw, ImageFont
    doc=ET.parse(xml_path).getroot(); paths=[]
    import os
    candidates = [Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'arial.ttf',
                  Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
                  Path('/System/Library/Fonts/Supplemental/Arial.ttf')]
    font_path = next((p for p in candidates if p.exists()), None)
    small = ImageFont.truetype(str(font_path),14) if font_path else ImageFont.load_default(size=14)
    for diagram in doc.findall('diagram'):
        img=Image.new('RGB',(1450,980),'white'); d=ImageDraw.Draw(img)
        cells={c.get('id'):c for c in diagram.findall('.//mxCell')}
        geom={uid:(float(c.find('mxGeometry').get('x')),float(c.find('mxGeometry').get('y')),
                   float(c.find('mxGeometry').get('width')),float(c.find('mxGeometry').get('height')))
              for uid,c in cells.items() if c.get('vertex')=='1' and c.find('mxGeometry') is not None}
        for c in cells.values():
            if c.get('edge')!='1':continue
            a=geom.get(c.get('source'));b=geom.get(c.get('target'))
            if not a or not b:continue
            x1=a[0]+a[2]; y1=a[1]+a[3]/2; x2=b[0];y2=b[1]+b[3]/2
            if x2<x1:x1=a[0]+a[2]/2;y1=a[1]+a[3];x2=b[0]+b[2]/2;y2=b[1]
            d.line((x1,y1,x2,y2),fill='#64748B',width=3)
            # Simple arrow tip.
            if x2>x1: d.polygon([(x2,y2),(x2-10,y2-5),(x2-10,y2+5)],fill='#64748B')
        for uid,(x,y,w,h) in geom.items():
            c=cells[uid]; style=c.get('style','')
            fill='#DBEAFE' if 'fillColor=#DBEAFE' in style else '#FFF7ED' if 'fillColor=#FFF7ED' in style else '#EFF6FF' if 'fillColor=#EFF6FF' in style else '#F3F6FB'
            d.rounded_rectangle((x,y,x+w,y+h),radius=10,outline='#94A3B8',fill=fill,width=2)
            text=c.get('value','').replace('&lt;','<')
            words=text.split();lines=[];cur=''
            for word in words:
                attempt=(cur+' '+word).strip()
                if d.textbbox((0,0),attempt,font=small)[2] < w-22: cur=attempt
                else:
                    if cur:lines.append(cur)
                    cur=word
            if cur:lines.append(cur)
            total=len(lines)*19
            yy=y+max(6,(h-total)/2)
            for line in lines[:min(8,max(1,int(h//19)))]:
                bb=d.textbbox((0,0),line,font=small)
                d.text((x+max(8,(w-(bb[2]-bb[0]))/2),yy),line,font=small,fill='#172554')
                yy+=19
        name=diagram.get('name','diagram').replace(' ','_')
        # Crop unused virtual draw.io canvas while retaining every editable shape.
        if geom:
            x0=max(0,int(min(r[0] for r in geom.values())-40))
            y0=max(0,int(min(r[1] for r in geom.values())-40))
            x1=min(img.width,int(max(r[0]+r[2] for r in geom.values())+40))
            y1=min(img.height,int(max(r[1]+r[3] for r in geom.values())+40))
            img=img.crop((x0,y0,x1,y1))
        p=Path(output_dir)/(name+'.png');p.parent.mkdir(parents=True,exist_ok=True);img.save(p)
        paths.append((diagram.get('name'),p))
    return paths


def make_docx(out, *, cases=None, run=None, diagram_path='course/drawio/CMPE187_All_Editable_DrawIO.drawio'):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    out=Path(out);out.parent.mkdir(parents=True,exist_ok=True)
    cases=cases or starter_cases();validate_cases(cases,require_full_matrix=True)
    if run and (run.get('schema') != 'mm-emotion.multi-model.v1' or run.get('mode') != 'matrix'):
        raise ValueError('Only genuine multi-model rubric matrix evidence is accepted')
    results=(run or {}).get('results',[])
    if any(r.get('origin') != 'real_model_inference' for r in results):
        raise ValueError('Report rejects synthetic or unverified experiment results')
    stats=summary(results)
    doc=Document();section=doc.sections[0];section.top_margin=Inches(.85);section.bottom_margin=Inches(.8);section.left_margin=Inches(.87);section.right_margin=Inches(.87)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6)
    for name,size,color in [('Title',19,RGBColor(16,42,80)),('Heading 1',14,RGBColor(16,42,80)),('Heading 2',11,RGBColor(32,72,114)),('Heading 3',10,RGBColor(32,72,114))]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    def p(text='',style=None):return doc.add_paragraph(text,style)
    def h(t,n=1):doc.add_heading(t,n)
    def tab(headers,rows,small=False):
        table=doc.add_table(rows=1,cols=len(headers));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,t in enumerate(headers):
            c=table.rows[0].cells[i];c.text=str(t);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in c.paragraphs:
                for run in par.runs:run.bold=True;run.font.size=Pt(8 if small else 9)
        for row in rows:
            cells=table.add_row().cells
            for i,t in enumerate(row):
                cells[i].text=str(t);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for par in cells[i].paragraphs:
                    for run in par.runs:run.font.size=Pt(7.5 if small else 8.5)
        doc.add_paragraph()
        return table
    doc.add_heading('AI Function Test Document',0)
    p('CMPE 187 | Deliverable #2 | Emotional Activation Vectors in Language Models')
    p('Project: Model Mountain Emotion Lab  |  Evidence: '+('Model inference run supplied' if results else 'Not yet executed (template and test models only)'))
    p('Editable report generated from the Deliverable #2A classification contract. Original thresholds remain documented; proposed interpretations are unapproved.')
    if results:
        tab(['Alias','Checkpoint','Precision','Revision prefix'],
            [[m.get('alias','?'),m.get('model_id','?'),m.get('precision','?'),str(m.get('resolved_revision') or 'unrecorded')[:12]]
             for m in run.get('models',[])],small=True)
        p('Full checkpoint revisions, prompts, seeds, durations, coefficients, vector hashes, runtime details and recorded failures accompany the JSON evidence. This report presents a small exploratory run, not calibrated emotion accuracy.')
    h('Section 1 – AI Function Test Requirement Analysis')
    p('This study explores whether model-internal activations contain directions correlated with happiness, sadness, and anger. It uses matched neutral-versus-emotional prompts and cosine similarity between held-out activation differences and vectors derived from disjoint training prompts. Text-only diction changes are not treated as proof of internal emotional understanding.')
    for idx,e in enumerate(EMOTIONS,1):p(f'R{idx}: Identify and compare {e}-related vectors across neutral, moderate, and high intensity settings. Evaluate stability across objective/stated/enacted abstraction and standard/expressive diction.')
    p('Functional requirements: permit editing/importing prompts, capture real model activations, fit baseline directions without test leakage, classify the held-out vector, evaluate ranges after threshold approval, compare multiple checkpoints, and export reproducible evidence. Unsupported models and failed captures must be reported, not silently simulated.')
    h('Section 2 – AI Test Modeling for Selected AI Features')
    previews=render_drawio_png(diagram_path,out.parent/'diagram_previews')
    titles=['2.1 Context Modeling for AI Features','2.2 AI Function Input Classifications','2.3 AI Function Output/Event/Action Classifications','2.4 AI Function Classification Decision Table']
    for i,t in enumerate(titles):
        h(t,2)
        if i==0:p('Context branches: abstraction (objective, stated, enacted) and diction (standard, expressive).')
        elif i==1:p('Input branches: happiness, sadness, anger, each with neutral, moderate and high intensity. Intensity is a derived tag in the 18-case minimum, not an independent fully crossed dimension.')
        elif i==2:p('Capture model output, calculate activation delta, compare with each trained emotion direction, check expected emotion, then check an approved range. Cases without reviewed ranges remain ungraded; zero-norm vectors are unscorable.')
        else:p('All three emotions share the following six context classifications; 3 x 6 = 18 required primary cases.')
        if i<3:
            doc.add_picture(str(previews[i][1]),width=Inches(6.25))
            cap=p('Editable source: course/drawio/'+previews[i][0].replace(' ','_')+'.drawio');cap.style='Caption'
        else:
            tab(['Abstraction','Diction','Intensity','Proposed range','PDF source issue'],
                [[a,d,cases[j]['intensity'],f'{cases[j]["proposed_min"]:.1f} to {cases[j]["proposed_max"]:.1f}',
                 'Reversed inequality' if j%6 else 'Consistent']
                 for j,(a,d) in enumerate(( (x,y) for x in ABSTRACTIONS for y in DICTIONS))],small=True)
            doc.add_picture(str(previews[i][1]),width=Inches(6.15))
    h('Section 3 – AI Function Test Cases with Inputs/Expected Outputs')
    h('3.1 Test data models',2)
    p('The full 18-case context matrix uses three emotions x three abstraction categories x two diction categories. Each case stores prompt, neutral reference, expected emotion, intensity, source range, proposed bounds and run metadata. Baseline-vector training uses twelve separate prompt pairs across distinct situations. Test prompts are not used in training.')
    h('3.2 Test case reports',2)
    if results:
        result_rows = []
        # Show every model/case rather than silently overwriting duplicate case IDs.
        for record in results:
            case = next((c for c in cases if c['case_id']==record.get('case_id')),None)
            if case is None: continue
            v=record.get('expected_cosine')
            result_rows.append([case['case_id'], str(record.get('model_alias') or record.get('model_id','?').split('/')[-1]),
                case['emotion'],case['abstraction']+'/'+case['diction'],case['intensity'],
                f"{float(case['proposed_min']):.2f} .. {float(case['proposed_max']):.2f}",
                f"{float(v):.3f}" if v is not None else 'Unscorable',
                {'ungraded_provisional_thresholds':'Ungraded','unscorable_vector':'Unscorable'}.get(record.get('status'),record.get('status','?'))])
    else:
        result_rows = [[case['case_id'],'Not run',case['emotion'],case['abstraction']+'/'+case['diction'],
            case['intensity'],f"{float(case['proposed_min']):.2f} .. {float(case['proposed_max']):.2f}",
            'Not run','Not run'] for case in cases]
    tab(['ID','Model','Emotion','Context','Intensity','Proposed cosine','Observed','Status'],result_rows,small=True)
    h('3.3 Test case analysis (statistics)',2)
    p('Minimum classification coverage: 3 emotions x 3 abstraction categories x 2 diction categories = 18 test cases per model. If intensity were treated as an independent factor, exhaustive coverage would be 54, but that is not the source document’s current design.')
    p(f"Executed records: {len(results)}. Scorable records: {stats['scorable_cases']}. Mean expected cosine: "+(f"{stats['mean_expected_cosine']:.4f}." if stats['mean_expected_cosine'] is not None else 'Not available.'))
    if results:tab(['Result status','Count'],sorted(stats['status_counts'].items()))
    else:p('No GPU-backed inference results were supplied. No pass rate, accuracy, confusion matrix, or robustness statistics can be reported yet.')
    h('Section 4 – AI Function Test Results and Test Criteria')
    h('4.1 AI test model-based test results analysis',2)
    if results:
        p('The table above reports observed classifications and similarities. Results are model-specific, and should only be compared under matched prompt sets, sampling parameters, layer selection protocol, and quantization configuration.')
        by_model=defaultdict(list)
        for r in results:by_model[str(r.get('model_id','unspecified'))+' / '+str(r.get('precision','?'))].append(r)
        tab(['Model and precision','Records','Graded pass','Fail','Ungraded/unscorable'],[[m,len(rs),sum(r.get('status')=='pass' for r in rs),sum(r.get('status')=='fail' for r in rs),sum(r.get('status') not in ('pass','fail') for r in rs)] for m,rs in by_model.items()],small=True)
    else:p('Status: NOT RUN. Conduct actual forward-pass captures on an authorized compatible model and import the generated JSON results. Synthetic numerical smoke-test outputs are not permitted as test evidence.')
    supplemental=(run or {}).get('supplemental',{})
    if supplemental:
        p('Additional measured intervention pilot: '+str(supplemental.get('steering_generations',0))+' generation records across three native checkpoints, with coefficients -3, 0 and +3 at sampling seeds 42 and 43. Vector discovery scanned blocks 5, 11 and 17; generation interventions used block 5. These are training-only exploratory vectors evaluated on a separate exam-advice prompt. Text changes require independent blinded labeling before interpreting an emotion effect.')
        comparison=supplemental.get('precision_comparison',{})
        if comparison.get('cosine_by_emotion_layer'):
            p('Matched weight-precision pilot: Qwen2.5 0.5B FP16 versus NF4 at block 5, using identical dataset/checkpoint and two pairs per emotion per split. Validation and held-out test splits remain separate from vector fitting.')
            tab(['Emotion','Block','FP16 versus NF4 direction cosine'],
                [[r['emotion'],r['layer_0based'],f"{r['direction_cosine']:.4f}"] for r in comparison['cosine_by_emotion_layer']],small=True)
            p('Representation similarity alone does not establish preserved behavior or a calibrated emotion construct. These four benchmark emotions differ from the course matrix labels. Qwen2.5 1.5B FP16 produced nonfinite activations in an earlier run; the successful 1.5B experiment explicitly uses FP32. The original failure evidence is retained.')
    p('Test complexity includes model loading and two hidden-state captures per case, plus contrastive training captures. Layer scanning multiplies capture and analysis cost; collect requested layers together where supported rather than reloading the model repeatedly.')
    h('4.2 AI Function bug analysis',2)
    tab(['ID','Issue / impact','Recommendation'],[
        ('REQ-01','Deliverable #2A pages 5–7 reverse lower/upper cosine bounds for five of six rows per emotion.','Obtain instructor approval for proposed ascending ranges before automated grading.'),
        ('REQ-02','Neutral cases still list the target emotion as expected; no-match vs correct-label semantics are ambiguous.','Clarify whether neutral should classify as neutral or simply have low similarity.'),
        ('REQ-03','Cosine thresholds and 0.9 close-match rule lack independent calibration.','Estimate thresholds from separate validation examples; reserve test split.'),
        ('METH-01','Changing wording may move activations for reasons unrelated to emotional abstraction.','Add topic-matched negative controls and paraphrases.'),
        ('METH-02','Near-zero activation differences have unstable cosine.','Mark unscorable; never substitute zero or silently pass.')],small=True)
    if results:
        failed=[r for r in results if r.get('status') in ('fail','unscorable_vector')]
        if failed:
            p('Observed failing/unscorable cases from model executions (these are measured outcome issues, not automatically implementation bugs):')
            tab(['Model','Case','Status','Predicted','Expected','Cosine'],
                [[str(r.get('model_id','?')).split('/')[-1],r.get('case_id','?'),r.get('status','?'),
                  str(r.get('predicted_emotion','?')),str(r.get('expected_emotion','?')),
                  f"{r.get('expected_cosine'):.3f}" if isinstance(r.get('expected_cosine'),(float,int)) else 'N/A']
                  for r in failed[:60]],small=True)
        else:
            p('No explicitly failed or unscorable runs were submitted. This is not evidence that the model has no errors: thresholds may still be unapproved.')
    h('4.3 AI function test quality assessment',2)
    p('Coverage quality: the 18-case classification matrix is complete when all rows execute for each selected model. Scientific validity requires independent baseline training, no prompt leakage, separate threshold calibration, repeated seeded runs, error/latency logging, and independent annotation of generated text or downstream behavior.')
    p('Acceptance criteria: all 18 unique combinations represented; every empirical result traceable to model/checkpoint, precision, layer, prompt, reference and seed; no fabricated evidence; failures recorded; reviewed scoring thresholds; editable draw.io artifacts and reproducible CSV/JSON output.')
    p('Limitations: an activation direction correlated with emotional language does not prove subjective experience or a uniquely identifiable emotion. The starter dataset is synthetic, small and English-language; conclusions remain exploratory.')
    p('Source: CMPE 187 Group Deliverable #2A, pages 1–8 (provided). Diagram previews are generated from accompanying editable draw.io files; prompt wording in this suite is newly authored.')
    h('Appendix Full Test Inputs and Expected Outputs')
    for case in cases:
        h(case['case_id']+' '+case['emotion'],2)
        p('Input: '+case['prompt'])
        p('Neutral reference: '+case['neutral_reference'])
        p('Expected label: '+case['expected_emotion']+'. Intensity: '+case['intensity']+'.')
        p('Original source range: '+case['source_range_text']+'. Proposed bounds remain unapproved; this case is ungraded until review.')
    doc.save(out)
    return out


def main():
    p=argparse.ArgumentParser(description='Generate CMPE187 Deliverable #2 DOCX from actual results')
    p.add_argument('--out',default='course/generated/CMPE187_Deliverable_2.docx')
    p.add_argument('--results',default=None)
    p.add_argument('--diagrams',default='course/drawio/CMPE187_All_Editable_DrawIO.drawio')
    args=p.parse_args()
    run=json.loads(Path(args.results).read_text(encoding='utf8')) if args.results else None
    print(make_docx(args.out,run=run,diagram_path=args.diagrams))
if __name__=='__main__':main()
